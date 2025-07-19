import os
import unittest
from datetime import datetime
from tempfile import NamedTemporaryFile
from typing import cast

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

# Исправленный импорт с игнорированием ошибки mypy
try:
    from HomeWork.fin_operation import (Transaction, TransactionState,  # type: ignore[import-not-found]
                                        get_transaction_stats, read_transactions_from_csv,
                                        read_transactions_from_excel)
except ImportError:
    from src.fin_operation import (Transaction, TransactionState,  # type: ignore[import-not-found]
                                   get_transaction_stats, read_transactions_from_csv, read_transactions_from_excel)


class TestTransaction(unittest.TestCase):
    def test_transaction_creation(self) -> None:
        """Тестирование создания объекта Transaction"""
        dt = datetime.now()
        t = Transaction(
            id=1,
            state=TransactionState.EXECUTED,
            date=dt,
            amount=100.0,
            currency_name="US Dollar",
            currency_code="USD",
            from_account="1234567890",
            to_account="0987654321",
            description="Payment"
        )

        self.assertEqual(t.id, 1)
        self.assertEqual(t.state, TransactionState.EXECUTED)
        self.assertEqual(t.date, dt)
        self.assertEqual(t.amount, 100.0)
        self.assertEqual(t.currency_name, "US Dollar")
        self.assertEqual(t.currency_code, "USD")
        self.assertEqual(t.from_account, "1234567890")
        self.assertEqual(t.to_account, "0987654321")
        self.assertEqual(t.description, "Payment")

    def test_transaction_without_from_account(self) -> None:
        """Тестирование создания Transaction без from_account"""
        t = Transaction(
            id=2,
            state=TransactionState.PENDING,
            date=datetime.now(),
            amount=50.0,
            currency_name="Euro",
            currency_code="EUR",
            from_account=None,
            to_account="1122334455",
            description="Deposit"
        )

        self.assertIsNone(t.from_account)


class TestCSVReader(unittest.TestCase):
    def setUp(self) -> None:
        """Создаем временный CSV файл для тестирования"""
        self.csv_content = """id;state;date;amount;currency_name;currency_code;from;to;description
1;EXECUTED;2023-01-01T12:00:00Z;100.0;US Dollar;USD;1234567890;0987654321;Payment
2;CANCELED;2023-01-02T13:00:00Z;200.0;Euro;EUR;;1122334455;Deposit
3;PENDING;2023-01-03T14:00:00Z;300.0;Pound;GBP;3344556677;8899001122;Transfer
"""

        self.empty_csv_content = """id;state;date;amount;currency_name;currency_code;from;to;description
"""

        self.invalid_csv_content = """id;state;date;amount;currency_name;currency_code;from;to;description
4;INVALID;2023-01-04T15:00:00Z;400.0;Yen;JPY;5566778899;0011223344;Invalid
5;EXECUTED;invalid-date;500.0;Yuan;CNY;6677889900;1122334455;Bad date
"""

        self.csv_file = NamedTemporaryFile(mode='w', delete=False, suffix='.csv', encoding='utf-8')
        self.csv_file.write(self.csv_content)
        self.csv_file.close()

        self.empty_csv_file = NamedTemporaryFile(mode='w', delete=False, suffix='.csv', encoding='utf-8')
        self.empty_csv_file.write(self.empty_csv_content)
        self.empty_csv_file.close()

        self.invalid_csv_file = NamedTemporaryFile(mode='w', delete=False, suffix='.csv', encoding='utf-8')
        self.invalid_csv_file.write(self.invalid_csv_content)
        self.invalid_csv_file.close()

    def tearDown(self) -> None:
        """Удаляем временные файлы после тестов"""
        os.unlink(self.csv_file.name)
        os.unlink(self.empty_csv_file.name)
        os.unlink(self.invalid_csv_file.name)

    def test_read_valid_csv(self) -> None:
        """Тестирование чтения валидного CSV файла"""
        transactions = read_transactions_from_csv(self.csv_file.name)

        self.assertEqual(len(transactions), 3)

        # Проверяем первую транзакцию
        self.assertEqual(transactions[0].id, 1)
        self.assertEqual(transactions[0].state, TransactionState.EXECUTED)
        self.assertEqual(transactions[0].date, datetime(2023, 1, 1, 12, 0, 0))
        self.assertEqual(transactions[0].amount, 100.0)
        self.assertEqual(transactions[0].currency_name, "US Dollar")
        self.assertEqual(transactions[0].currency_code, "USD")
        self.assertEqual(transactions[0].from_account, "1234567890")
        self.assertEqual(transactions[0].to_account, "0987654321")
        self.assertEqual(transactions[0].description, "Payment")

        # Проверяем вторую транзакцию (без from_account)
        self.assertIsNone(transactions[1].from_account)

    def test_read_empty_csv(self) -> None:
        """Тестирование чтения пустого CSV файла"""
        transactions = read_transactions_from_csv(self.empty_csv_file.name)
        self.assertEqual(len(transactions), 0)

    def test_read_invalid_csv(self) -> None:
        """Тестирование чтения CSV с невалидными данными"""
        transactions = read_transactions_from_csv(self.invalid_csv_file.name)
        self.assertEqual(len(transactions), 0)


class TestExcelReader(unittest.TestCase):
    def setUp(self) -> None:
        """Создаем временный Excel файл для тестирования"""
        # Создаем валидный Excel файл
        self.valid_excel_file = NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb = Workbook()
        ws = cast(Worksheet, wb.active)  # Явное приведение типа
        ws.title = "Transactions"

        # Заголовки
        ws.append([
            'id', 'state', 'date', 'amount', 'currency_name',
            'currency_code', 'from', 'to', 'description'
        ])

        # Данные
        ws.append([
            1, 'EXECUTED', '2023-01-01T12:00:00Z', 100.0,
            'US Dollar', 'USD', '1234567890', '0987654321', 'Payment'
        ])
        ws.append([
            2, 'CANCELED', '2023-01-02T13:00:00Z', 200.0,
            'Euro', 'EUR', None, '1122334455', 'Deposit'
        ])
        ws.append([
            3, 'PENDING', '2023-01-03T14:00:00Z', 300.0,
            'Pound', 'GBP', '3344556677', '8899001122', 'Transfer'
        ])

        wb.save(self.valid_excel_file.name)
        wb.close()

        # Создаем пустой Excel файл
        self.empty_excel_file = NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = "Empty"
        ws.append([
            'id', 'state', 'date', 'amount', 'currency_name',
            'currency_code', 'from', 'to', 'description'
        ])
        wb.save(self.empty_excel_file.name)
        wb.close()

        # Создаем Excel файл с невалидными данными
        self.invalid_excel_file = NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.title = "Invalid"

        ws.append([
            'id', 'state', 'date', 'amount', 'currency_name',
            'currency_code', 'from', 'to', 'description'
        ])
        ws.append([
            4, 'INVALID', '2023-01-04T15:00:00Z', 400.0,
            'Yen', 'JPY', '5566778899', '0011223344', 'Invalid'
        ])
        ws.append([
            5, 'EXECUTED', 'invalid-date', 500.0,
            'Yuan', 'CNY', '6677889900', '1122334455', 'Bad date'
        ])

        wb.save(self.invalid_excel_file.name)
        wb.close()

    def tearDown(self) -> None:
        """Удаляем временные файлы после тестов"""
        for filename in [
            self.valid_excel_file.name,
            self.empty_excel_file.name,
            self.invalid_excel_file.name
        ]:
            try:
                if os.path.exists(filename):
                    os.unlink(filename)
            except PermissionError:
                pass

    def test_read_valid_excel(self) -> None:
        """Тестирование чтения валидного Excel файла"""
        transactions = read_transactions_from_excel(self.valid_excel_file.name, "Transactions")

        self.assertEqual(len(transactions), 3)

        # Проверяем первую транзакцию
        self.assertEqual(transactions[0].id, 1)
        self.assertEqual(transactions[0].state, TransactionState.EXECUTED)
        self.assertEqual(transactions[0].date, datetime(2023, 1, 1, 12, 0, 0))
        self.assertEqual(transactions[0].amount, 100.0)
        self.assertEqual(transactions[0].currency_name, "US Dollar")
        self.assertEqual(transactions[0].currency_code, "USD")
        self.assertEqual(transactions[0].from_account, "1234567890")
        self.assertEqual(transactions[0].to_account, "0987654321")
        self.assertEqual(transactions[0].description, "Payment")

        # Проверяем вторую транзакцию (без from_account)
        self.assertIsNone(transactions[1].from_account)

    def test_read_empty_excel(self) -> None:
        """Тестирование чтения пустого Excel файла"""
        transactions = read_transactions_from_excel(self.empty_excel_file.name, "Empty")
        self.assertEqual(len(transactions), 0)

    def test_read_invalid_excel(self) -> None:
        """Тестирование чтения Excel с невалидными данными"""
        transactions = read_transactions_from_excel(self.invalid_excel_file.name, "Invalid")
        self.assertEqual(len(transactions), 0)

    def test_read_default_sheet(self) -> None:
        """Тестирование чтения активного листа по умолчанию"""
        transactions = read_transactions_from_excel(self.valid_excel_file.name)
        self.assertEqual(len(transactions), 3)


class TestTransactionStats(unittest.TestCase):
    def test_stats_with_transactions(self) -> None:
        """Тестирование статистики с транзакциями"""
        transactions = [
            Transaction(
                id=1,
                state=TransactionState.EXECUTED,
                date=datetime(2023, 1, 1),
                amount=100.0,
                currency_name="USD",
                currency_code="USD",
                from_account="123",
                to_account="456",
                description="Test"
            ),
            Transaction(
                id=2,
                state=TransactionState.EXECUTED,
                date=datetime(2023, 1, 2),
                amount=200.0,
                currency_name="EUR",
                currency_code="EUR",
                from_account="456",
                to_account="789",
                description="Test"
            ),
            Transaction(
                id=3,
                state=TransactionState.CANCELED,
                date=datetime(2023, 1, 3),
                amount=300.0,
                currency_name="GBP",
                currency_code="GBP",
                from_account="789",
                to_account="012",
                description="Test"
            ),
            Transaction(
                id=4,
                state=TransactionState.PENDING,
                date=datetime(2023, 1, 4),
                amount=400.0,
                currency_name="USD",
                currency_code="USD",
                from_account="012",
                to_account="345",
                description="Test"
            )
        ]

        stats = get_transaction_stats(transactions)

        self.assertEqual(stats['total_transactions'], 4)
        self.assertEqual(stats['executed_count'], 2)
        self.assertEqual(stats['canceled_count'], 1)
        self.assertEqual(stats['pending_count'], 1)
        self.assertEqual(stats['total_amount'], 1000.0)
        self.assertEqual(stats['currencies'], {'USD', 'EUR', 'GBP'})
        self.assertEqual(stats['first_date'], datetime(2023, 1, 1))
        self.assertEqual(stats['last_date'], datetime(2023, 1, 4))

    def test_stats_with_empty_list(self) -> None:
        """Тестирование статистики с пустым списком транзакций"""
        stats = get_transaction_stats([])
        self.assertEqual(stats, {})


if __name__ == '__main__':
    unittest.main()
